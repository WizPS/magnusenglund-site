"""Convert the Helsingborg election workbook into a compact JSON data file.

Usage:
    python scripts/convert-valanalys.py input.xlsx app/data/val-2026-helsingborg.json

The source workbook is not needed by the website at runtime. Only the generated
JSON is published with the static site.
"""

from __future__ import annotations

import argparse
import json
from collections import OrderedDict
from pathlib import Path
from typing import Any

import openpyxl


def text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def number(value: Any) -> int:
    if value is None or value == "":
        return 0
    if isinstance(value, bool):
        return int(value)
    return int(round(float(value)))


def convert(source: Path, destination: Path) -> None:
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    if "Personroster" not in workbook.sheetnames:
        raise ValueError("The workbook does not contain a Personroster sheet")

    sheet = workbook["Personroster"]
    rows = sheet.iter_rows(values_only=True)
    headers = [text(value) for value in next(rows)]
    indexes = {header: index for index, header in enumerate(headers)}

    required = {
        "namn",
        "valdistriktskod",
        "antalRoster.1",
        "partibeteckning",
        "partiforkortning",
        "kandidat",
        "kandidatnummer",
        "antalPersonroster",
    }
    missing = sorted(required - indexes.keys())
    if missing:
        raise ValueError(f"Missing required columns: {', '.join(missing)}")

    parties: OrderedDict[str, dict[str, Any]] = OrderedDict()
    seen_party_districts: set[tuple[str, str]] = set()

    def value(row: tuple[Any, ...], key: str) -> Any:
        return row[indexes[key]]

    for row in rows:
        party_name = text(value(row, "partibeteckning"))
        if not party_name:
            continue

        party_code = text(value(row, "partiforkortning")) or None
        party_key = party_code or party_name
        district_name = text(value(row, "namn")) or "Okänt valdistrikt"
        district_code = text(value(row, "valdistriktskod")) or None
        district_key = district_code or district_name

        party = parties.setdefault(
            party_key,
            {
                "code": party_code or party_name,
                "name": party_name,
                "partyVotes": 0,
                "candidates": OrderedDict(),
                "districts": OrderedDict(),
            },
        )

        district = party["districts"].setdefault(
            district_key,
            {
                "code": district_code,
                "name": district_name,
                "partyVotes": 0,
            },
        )

        pair_key = (party_key, district_key)
        if pair_key not in seen_party_districts:
            party_votes = number(value(row, "antalRoster.1"))
            district["partyVotes"] += party_votes
            party["partyVotes"] += party_votes
            seen_party_districts.add(pair_key)

        candidate_name = text(value(row, "kandidat"))
        if not candidate_name:
            continue

        candidate_number = value(row, "kandidatnummer")
        candidate_key = text(candidate_number) or candidate_name
        candidate = party["candidates"].setdefault(
            candidate_key,
            {
                "number": number(candidate_number) if candidate_number not in (None, "") else None,
                "name": candidate_name,
                "personalVotes": 0,
                "districts": OrderedDict(),
            },
        )
        personal_votes = number(value(row, "antalPersonroster"))
        candidate["personalVotes"] += personal_votes
        candidate_district = candidate["districts"].setdefault(
            district_key,
            {
                "code": district_code,
                "name": district_name,
                "votes": 0,
            },
        )
        candidate_district["votes"] += personal_votes

    party_values = []
    for party in parties.values():
        candidates = []
        for candidate in party["candidates"].values():
            candidate["districts"] = sorted(
                (district for district in candidate["districts"].values() if district["votes"] > 0),
                key=lambda district: (-district["votes"], district["name"]),
            )
            candidates.append(candidate)
        party["candidates"] = sorted(
            candidates,
            key=lambda candidate: (-candidate["personalVotes"], candidate["name"]),
        )
        party["districts"] = sorted(
            party["districts"].values(),
            key=lambda district: district["name"],
        )
        party_values.append(party)

    party_values.sort(key=lambda party: party["name"])
    payload = {
        "source": {
            "name": "Valmyndigheten – Valanalys ME.xlsx",
            "sheet": "Personroster",
            "note": "Bearbetad från slutlig rösträkning för Helsingborgs kommun.",
        },
        "municipality": "Helsingborg",
        "municipalityCode": "1283",
        "election": "2026",
        "electionType": "KF",
        "parties": party_values,
    }

    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    print(f"Wrote {destination} with {len(party_values)} parties")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("destination", type=Path)
    args = parser.parse_args()
    convert(args.source, args.destination)


if __name__ == "__main__":
    main()
